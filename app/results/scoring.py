"""Чтение выгрузки Яндекс.Формы и подсчёт баллов по ключу.

Здесь только чистые функции без базы — их удобно проверять тестами.
Команды, которые пишут в базу, — в app/cli/qual.py.

Что ожидается от выгрузки (XLSX или CSV):
- первая строка — заголовки;
- столбец с кодом участника: в заголовке есть «код участника»;
- столбцы задач: заголовок начинается с номера, например «1. Сколько…»,
  «Задача 3», «№5»;
- если есть столбец со временем ответа («время», «дата»), повторные
  ответы одного участника определяются по нему, иначе по порядку строк.
  Засчитывается ПЕРВЫЙ ответ.
"""

import csv
import io
import re
from dataclasses import dataclass, field
from datetime import datetime
from decimal import Decimal
from pathlib import Path

TASK_HEADER_RE = re.compile(r"^\s*(?:задача|задание|вопрос)?\s*№?\s*(\d+)\b", re.IGNORECASE)
TIME_FORMATS = ("%Y-%m-%d %H:%M:%S", "%d.%m.%Y %H:%M:%S", "%Y-%m-%d %H:%M", "%d.%m.%Y %H:%M")


class TableError(Exception):
    pass


def read_table(path: Path) -> list[list]:
    if path.suffix.lower() in (".xlsx", ".xlsm"):
        from openpyxl import load_workbook

        sheet = load_workbook(path, read_only=True, data_only=True).active
        return [list(row) for row in sheet.iter_rows(values_only=True)]
    text = path.read_bytes().decode("utf-8-sig")
    dialect = csv.Sniffer().sniff(text[:4096], delimiters=",;\t")
    return [row for row in csv.reader(io.StringIO(text), dialect)]


def cell_text(value) -> str:
    if value is None:
        return ""
    if isinstance(value, float) and value.is_integer():
        value = int(value)
    return str(value).strip()


def normalize_answer(value) -> str:
    text = cell_text(value).lower().replace("ё", "е")
    text = " ".join(text.split())
    if re.fullmatch(r"-?\d+,\d+", text):
        text = text.replace(",", ".")
    return text


def normalize_code(value) -> str:
    return re.sub(r"[\s\-]", "", cell_text(value)).upper()


def is_correct(answer, key_answer: str) -> bool:
    given = normalize_answer(answer)
    return bool(given) and given in {normalize_answer(v) for v in key_answer.split("|")}


@dataclass
class Layout:
    code_col: int
    task_cols: dict[int, int]  # номер задачи -> индекс столбца
    time_col: int | None


def detect_layout(headers: list) -> Layout:
    names = [cell_text(h) for h in headers]
    code_cols = [i for i, h in enumerate(names) if "код участника" in h.lower()]
    if len(code_cols) != 1:
        raise TableError(
            "Не найден (или найдено несколько) столбец «Код участника». Заголовки: "
            + " | ".join(names)
        )
    task_cols = {}
    for i, h in enumerate(names):
        if i == code_cols[0]:
            continue
        match = TASK_HEADER_RE.match(h)
        if match:
            task_no = int(match.group(1))
            if task_no in task_cols:
                raise TableError(f"Два столбца для задачи {task_no}")
            task_cols[task_no] = i
    if not task_cols:
        raise TableError(
            "Не найдены столбцы задач (заголовок должен начинаться с номера). Заголовки: "
            + " | ".join(names)
        )
    time_col = next(
        (i for i, h in enumerate(names) if re.search(r"время|дата", h, re.IGNORECASE)), None
    )
    return Layout(code_cols[0], task_cols, time_col)


def _parse_time(value) -> datetime | None:
    if isinstance(value, datetime):
        return value
    text = cell_text(value)
    for fmt in TIME_FORMATS:
        try:
            return datetime.strptime(text, fmt)
        except ValueError:
            continue
    return None


@dataclass
class Submission:
    code: str
    answers: dict[int, str]  # номер задачи -> ответ как в выгрузке


@dataclass
class Parsed:
    layout: Layout
    submissions: list[Submission] = field(default_factory=list)
    duplicates: dict[str, int] = field(default_factory=dict)  # код -> сколько лишних
    empty_code_rows: int = 0


def parse_submissions(rows: list[list]) -> Parsed:
    if not rows:
        raise TableError("Файл пустой")
    layout = detect_layout(rows[0])
    data = [r for r in rows[1:] if any(cell_text(c) for c in r)]

    if layout.time_col is not None:
        times = [_parse_time(r[layout.time_col]) if layout.time_col < len(r) else None for r in data]
        if all(t is not None for t in times):
            data = [r for _, r in sorted(zip(times, data, strict=True), key=lambda p: p[0])]

    parsed = Parsed(layout)
    seen: set[str] = set()
    for row in data:
        code = normalize_code(row[layout.code_col]) if layout.code_col < len(row) else ""
        if not code:
            parsed.empty_code_rows += 1
            continue
        if code in seen:
            parsed.duplicates[code] = parsed.duplicates.get(code, 0) + 1
            continue
        seen.add(code)
        answers = {
            task: cell_text(row[col]) if col < len(row) else ""
            for task, col in layout.task_cols.items()
        }
        parsed.submissions.append(Submission(code, answers))
    return parsed


@dataclass
class KeyItem:
    answer: str
    points: Decimal


def score(answers: dict[int, str], key: dict[int, KeyItem]) -> dict[int, Decimal]:
    """Балл за каждую задачу ключа; нет ответа или неверно — 0."""
    return {
        task: item.points if is_correct(answers.get(task, ""), item.answer) else Decimal(0)
        for task, item in key.items()
    }


def read_key(path: Path) -> dict[int, dict[int, KeyItem]]:
    """Ключ: таблица со столбцами класс, задача, ответ, баллы (первая
    строка — заголовки, названия любые). Возвращает {класс: {задача: …}}."""

    rows = read_table(path)
    key: dict[int, dict[int, KeyItem]] = {}
    for n, row in enumerate(rows[1:], start=2):
        if not any(cell_text(c) for c in row):
            continue
        try:
            grade, task = int(cell_text(row[0])), int(cell_text(row[1]))
            answer = cell_text(row[2])
            points = Decimal(cell_text(row[3]).replace(",", ".") or "1")
        except (ValueError, IndexError, ArithmeticError):
            raise TableError(f"Ключ, строка {n}: ожидается «класс; задача; ответ; баллы»") from None
        if not answer:
            raise TableError(f"Ключ, строка {n}: пустой ответ")
        if task in key.setdefault(grade, {}):
            raise TableError(f"Ключ, строка {n}: задача {task} для {grade} класса уже есть")
        key[grade][task] = KeyItem(answer, points)
    if not key:
        raise TableError("Ключ пустой")
    return key
